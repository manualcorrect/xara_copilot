import openpyxl
from xar_dom_engine import XarDocument

xlsx_path = r"C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\bahan\New folder\Aug\Template_Pekerjaan_Xara_Aug.xlsx"
xar_path = r"C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\bahan\New folder\Aug\0.xar"

wb = openpyxl.load_workbook(xlsx_path, data_only=True)
ws_m = wb['Tabel_Mutasi']

print("=== EXCEL TRANSACTIONS IN AUG ===")
tx_excel = []
for r in range(6, 50):
    val_nom = ws_m.cell(r, 5).value
    val_saldo = ws_m.cell(r, 7).value
    if val_nom is not None and val_saldo is not None:
        val_no = ws_m.cell(r, 1).value
        val_tgl = ws_m.cell(r, 2).value
        val_jam = ws_m.cell(r, 3).value
        tx_excel.append({
            "excel_row": r,
            "no": len(tx_excel) + 1,
            "orig_no": val_no,
            "tanggal": str(val_tgl or ''),
            "jam": str(val_jam or ''),
            "nominal": val_nom,
            "saldo": val_saldo
        })

print(f"Total transactions in Excel: {len(tx_excel)}")
for t in tx_excel:
    print(f"Tx {t['no']:02d} (Excel R{t['excel_row']}): Nom={t['nominal']:12} | Saldo={t['saldo']:12} | Tgl={t['tanggal']:15s} | Jam={t['jam']}")

print("\n=== 0.XAR IN AUG ===")
doc = XarDocument(xar_path)
print(f"Total records in 0.xar: {len(doc.records)}")
