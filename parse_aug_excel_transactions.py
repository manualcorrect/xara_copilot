import openpyxl
import json
import datetime

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\aug\Template_Pekerjaan_Xara_Aug.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws_tx = wb['Tabel_Mutasi']

tx_list = []
last_valid_date = "01 Aug 2026"

for r in range(6, 116): # rows 6 to 115 (110 transactions)
    row_no = ws_tx.cell(r, 1).value
    tgl_val = ws_tx.cell(r, 2).value
    jam_val = ws_tx.cell(r, 3).value
    uraian = ws_tx.cell(r, 4).value
    nom_val = ws_tx.cell(r, 5).value
    tipe_val = ws_tx.cell(r, 6).value
    saldo_val = ws_tx.cell(r, 7).value
    
    # Parse date
    date_str = ""
    if isinstance(tgl_val, datetime.datetime):
        date_str = tgl_val.strftime('%d Aug %Y')
    elif isinstance(tgl_val, str) and tgl_val.strip():
        s = tgl_val.strip()
        if '/' in s:
            parts = s.split('/')
            day = int(parts[0])
            date_str = f"{day:02d} Aug 2026"
        elif 'Aug' in s or '2026' in s:
            date_str = s
        else:
            date_str = s
    
    if date_str and ('Aug' in date_str or '2026' in date_str or '/' in str(tgl_val)):
        last_valid_date = date_str
    
    # Parse nominal
    nom_float = float(nom_val) if nom_val is not None else 0.0
    nom_type = "CR" if nom_float > 0 else "DB"
    
    # Format nominal: "+1.500.000,00" or "-3.500,00"
    abs_nom = abs(nom_float)
    int_part = int(abs_nom)
    dec_part = int(round((abs_nom - int_part) * 100))
    nom_str_formatted = f"{int_part:,}".replace(',', '.') + f",{dec_part:02d}"
    if nom_float > 0:
        formatted_nominal = f"+{nom_str_formatted}"
    else:
        formatted_nominal = f"-{nom_str_formatted}"
        
    # Format saldo: "1.521.347,81"
    saldo_float = float(saldo_val) if saldo_val is not None else 0.0
    s_int = int(saldo_float)
    s_dec = int(round((saldo_float - s_int) * 100))
    formatted_saldo = f"{s_int:,}".replace(',', '.') + f",{s_dec:02d}"
    
    tx_list.append({
        'row_idx': len(tx_list) + 1,
        'excel_row': r,
        'row_no': row_no,
        'raw_date': str(tgl_val),
        'resolved_date': last_valid_date,
        'raw_time': str(jam_val),
        'nom_float': nom_float,
        'nom_type': nom_type,
        'formatted_nominal': formatted_nominal,
        'saldo_float': saldo_float,
        'formatted_saldo': formatted_saldo,
        'uraian': str(uraian)
    })

print(f"[*] Parsed {len(tx_list)} transactions from Excel:")
for t in tx_list[:10]:
    print(f"  Row {t['row_idx']:2d}: {t['resolved_date']} | Nom: {t['formatted_nominal']} | Saldo: {t['formatted_saldo']}")
print(f"  ... (last 5 rows) ...")
for t in tx_list[-5:]:
    print(f"  Row {t['row_idx']:2d}: {t['resolved_date']} | Nom: {t['formatted_nominal']} | Saldo: {t['formatted_saldo']}")

with open('aug_parsed_tx.json', 'w', encoding='utf-8') as f:
    json.dump(tx_list, f, indent=2)
