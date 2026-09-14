import openpyxl

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\aug\Template_Pekerjaan_Xara_Aug.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Sheet 1: Header & Ringkasan
ws_info = wb['Header & Ringkasan']
print("=== INFORMASI NASABAH & RINGKASAN KEBERADAAN ===")
for r in range(1, ws_info.max_row + 1):
    vals = [ws_info.cell(r, c).value for c in range(1, 10)]
    if any(v is not None for v in vals):
        print(f"Row {r}: {vals}")

# 2. Sheet 2: Tabel_Mutasi
ws_tx = wb['Tabel_Mutasi']
print(f"\n=== TABEL MUTASI (Total Rows={ws_tx.max_row}) ===")
tx_rows = []
for r in range(4, ws_tx.max_row + 1):
    no = ws_tx.cell(r, 1).value
    tgl = ws_tx.cell(r, 2).value
    jam = ws_tx.cell(r, 3).value
    uraian = ws_tx.cell(r, 4).value
    nom = ws_tx.cell(r, 5).value
    tipe = ws_tx.cell(r, 6).value
    saldo = ws_tx.cell(r, 7).value
    if no is not None and str(no).strip() != '' and str(no).strip() != 'No':
        tx_rows.append({
            'excel_row': r,
            'no': no,
            'tgl': tgl,
            'jam': jam,
            'uraian': uraian,
            'nom': nom,
            'tipe': tipe,
            'saldo': saldo
        })

print(f"[*] Found {len(tx_rows)} transaction rows in Tabel_Mutasi:")
for t in tx_rows[:10]:
    print(f"  {t}")
print(f"  ... (last 5 rows) ...")
for t in tx_rows[-5:]:
    print(f"  {t}")
