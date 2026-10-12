import openpyxl
import os
import sys

def parse_xara_excel_template(excel_path):
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"File Excel tidak ditemukan: {excel_path}")
        
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # 1. Parse Sheet 1: Header & Ringkasan
    ws1 = wb["Header & Ringkasan"]
    config = {
        "project": {},
        "header": {},
        "pages": {},
        "dates": {},
        "summary": {},
        "transactions": [],
        "additional_adjustments": {
            "ruler": {},
            "colors": {},
            "text_overrides": [],
            "safety": {}
        }
    }
    
    # Helper for currency formatting
    def format_idr_currency(val):
        if val is None:
            return ""
        if isinstance(val, (int, float)):
            return f"{val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return str(val).strip()

    current_sec = None
    for r in range(5, ws1.max_row + 1):
        p_name = ws1.cell(row=r, column=1).value
        p_val = ws1.cell(row=r, column=2).value
        
        if not p_name:
            continue
        p_name = str(p_name).strip()
        p_val_str = str(p_val).strip() if p_val is not None else ""
        
        if p_name.startswith("1. INFORMASI FILE"):
            current_sec = "file"
        elif p_name.startswith("2. DATA HEADER"):
            current_sec = "header"
        elif p_name.startswith("3. PENOMORAN"):
            current_sec = "pages"
        elif p_name.startswith("4. BULAN"):
            current_sec = "dates"
        elif p_name.startswith("5. RINGKASAN"):
            current_sec = "summary"
        else:
            if current_sec == "file":
                if "File Sumber" in p_name: config["project"]["source_xar"] = p_val_str
                elif "File Output" in p_name: config["project"]["output_xar"] = p_val_str
                elif "Nama Dokumen" in p_name: config["project"]["job_title"] = p_val_str
            elif current_sec == "header":
                if "Nama Nasabah" in p_name: config["header"]["nama"] = p_val_str
                elif "Periode" in p_name: config["header"]["periode"] = p_val_str
                elif "Dicetak" in p_name: config["header"]["dicetak_pada"] = p_val_str
                elif "Rekening" in p_name: config["header"]["nomor_rekening"] = p_val_str
            elif current_sec == "pages":
                if "Mode Total Halaman" in p_name: config["pages"]["mode"] = p_val_str
                elif "Total Halaman Manual" in p_name: config["pages"]["manual_pages"] = p_val_str
            elif current_sec == "dates":
                if "Target Bulan" in p_name: config["dates"]["target_month_year"] = p_val_str
                elif "Digit Penutup" in p_name: config["dates"]["target_year_digit"] = p_val_str
            elif current_sec == "summary":
                formatted_summary = format_idr_currency(p_val)
                if "Saldo Awal" in p_name: config["summary"]["saldo_awal"] = formatted_summary
                elif "Dana Masuk" in p_name: config["summary"]["dana_masuk"] = formatted_summary
                elif "Dana Keluar" in p_name: config["summary"]["dana_keluar"] = formatted_summary
                elif "Saldo Akhir" in p_name: config["summary"]["saldo_akhir"] = formatted_summary

    # Auto-synchronize Periode with Target Bulan & Tahun if mismatched
    target_my = config["dates"].get("target_month_year", "").strip()
    periode_raw = config["header"].get("periode", "").strip()
    if target_my and periode_raw:
        t_parts = target_my.split()
        if len(t_parts) == 2:
            t_mon, t_yr = t_parts[0], t_parts[1]
            if t_mon.lower() not in periode_raw.lower():
                import calendar
                month_to_num = {'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'mei': 5, 'may': 5,
                                'jun': 6, 'jul': 7, 'agt': 8, 'aug': 8, 'sep': 9, 'okt': 10,
                                'oct': 10, 'nov': 11, 'des': 12, 'dec': 12}
                m_num = month_to_num.get(t_mon.lower(), 1)
                try:
                    yr_int = int(t_yr)
                    max_d = calendar.monthrange(yr_int, m_num)[1]
                    auto_per = f"01 {t_mon} {t_yr} - {max_d:02d} {t_mon} {t_yr}"
                    config["header"]["periode"] = auto_per
                except:
                    pass
                
    # 2. Parse Sheet 2: Tabel_Mutasi
    ws2 = wb["Tabel_Mutasi"]
    for r in range(5, ws2.max_row + 1):
        no_val = ws2.cell(row=r, column=1).value
        if str(no_val).strip() == '[AWAL]':
            continue

        tgl_raw = ws2.cell(row=r, column=2).value
        jam_raw = ws2.cell(row=r, column=3).value
        uraian_raw = ws2.cell(row=r, column=4).value
        nominal_raw = ws2.cell(row=r, column=5).value
        tipe_raw = ws2.cell(row=r, column=6).value
        saldo_raw = ws2.cell(row=r, column=7).value
        custom_color = str(ws2.cell(row=r, column=8).value or "AUTO").strip()
        custom_align = str(ws2.cell(row=r, column=9).value or "AUTO").strip()
        catatan = str(ws2.cell(row=r, column=10).value or "").strip()

        # Check if entire row is empty
        if no_val is None and tgl_raw is None and nominal_raw is None and saldo_raw is None and uraian_raw is None:
            continue

        # Sequential / Int numbering
        if no_val is not None:
            try:
                no_int = int(no_val)
            except:
                no_int = len(config["transactions"]) + 1
        else:
            no_int = len(config["transactions"]) + 1

        # Format nominal & saldo
        if isinstance(nominal_raw, (int, float)):
            nominal = format_idr_currency(nominal_raw)
        else:
            nominal = str(nominal_raw or "").strip()

        if isinstance(saldo_raw, (int, float)):
            saldo = format_idr_currency(saldo_raw)
        else:
            saldo = str(saldo_raw or "").strip()

        tipe = str(tipe_raw or "").strip().upper()
        if not tipe:
            if isinstance(nominal_raw, (int, float)):
                tipe = "CR" if nominal_raw > 0 else "DB"
            elif nominal.startswith("+"):
                tipe = "CR"
            elif nominal.startswith("-"):
                tipe = "DB"

        # Date & Time formatting
        from datetime import datetime as dt_type
        if isinstance(tgl_raw, dt_type):
            tgl = tgl_raw.strftime("%d %b %Y")
        else:
            tgl = str(tgl_raw or "").strip()
        if tgl and '/' not in tgl and '-' not in tgl and ' ' not in tgl:
            if tgl.isdigit() and int(tgl) > 31:
                tgl = ""

        jam = str(jam_raw or "").strip()
        uraian = str(uraian_raw or "").strip()

        # Check if nominal is zero
        nom_clean = nominal.replace("+", "").replace("-", "").replace(".", "").replace(",", ".").replace(" ", "").strip()
        is_zero = False
        try:
            val_float = float(nom_clean) if nom_clean else 0.0
            if abs(val_float) < 0.0001:
                is_zero = True
        except:
            if nom_clean in ("0", "0.00", "0,00", "-", ""):
                is_zero = True

        # Smart Normalization Rule: Skip zero-nominal transactions
        if is_zero and (uraian == "" or uraian == "-" or "0" in nominal):
            # Record filtered out per SOP Section IV
            continue

        config["transactions"].append({
            "no": len(config["transactions"]) + 1, # Sequential Reindexing (1..N)
            "orig_no": no_int,                      # Original number from Excel
            "tanggal": tgl,
            "jam": jam,
            "uraian": uraian,
            "nominal": nominal,
            "tipe": tipe,
            "saldo": saldo,
            "custom_color": custom_color,
            "custom_align": custom_align,
            "catatan": catatan
        })
        
        
    # 3. Parse Sheet 3: Penyesuaian_Tambahan
    ws3 = wb["Penyesuaian_Tambahan"]
    current_sec3 = None
    for r in range(5, ws3.max_row + 1):
        kat = ws3.cell(row=r, column=1).value
        param = ws3.cell(row=r, column=2).value
        val = ws3.cell(row=r, column=3).value
        target = ws3.cell(row=r, column=4).value
        ket = ws3.cell(row=r, column=5).value
        
        if not kat:
            continue
        kat = str(kat).strip()
        param = str(param).strip() if param is not None else ""
        val = str(val).strip() if val is not None else ""
        
        if kat.startswith("A. KUSTOMISASI RULER"):
            current_sec3 = "ruler"
        elif kat.startswith("B. KUSTOMISASI PALET"):
            current_sec3 = "colors"
        elif kat.startswith("C. OVERRIDE TEKS"):
            current_sec3 = "text_overrides"
        elif kat.startswith("D. FITUR KEAMANAN"):
            current_sec3 = "safety"
        else:
            if current_sec3 == "ruler":
                config["additional_adjustments"]["ruler"][param] = val
            elif current_sec3 == "colors":
                config["additional_adjustments"]["colors"][param] = val
            elif current_sec3 == "text_overrides":
                if param and not param.startswith("["):
                    config["additional_adjustments"]["text_overrides"].append({
                        "asal": param,
                        "baru": val,
                        "target": target,
                        "keterangan": ket
                    })
            elif current_sec3 == "safety":
                config["additional_adjustments"]["safety"][param] = val
                
    return config

if __name__ == "__main__":
    path = r"C:\Users\Lenovo\xara_copilot\Template_Pekerjaan_Xara.xlsx"
    cfg = parse_xara_excel_template(path)
    print("=========================================================")
    print("   PARSER TEST: TEMPLATE PEKERJAAN XARA (.XLSX) SUCCESS   ")
    print("=========================================================")
    print(f"[*] Proyek       : {cfg['project']['job_title']}")
    print(f"[*] Input File   : {cfg['project']['source_xar']}")
    print(f"[*] Output File  : {cfg['project']['output_xar']}")
    print(f"[*] Nasabah      : {cfg['header']['nama']}")
    print(f"[*] Periode      : {cfg['header']['periode']}")
    print(f"[*] Total Baris  : {len(cfg['transactions'])} Baris Transaksi Terbaca")
    print(f"[*] Baris 1      : {cfg['transactions'][0]['nominal']} ({cfg['transactions'][0]['tipe']}) -> Saldo: {cfg['transactions'][0]['saldo']}")
    print(f"[*] Baris 47     : {cfg['transactions'][46]['nominal']} ({cfg['transactions'][46]['tipe']}) -> Saldo: {cfg['transactions'][46]['saldo']}")
    print(f"[*] Saldo Akhir  : {cfg['summary']['saldo_akhir']}")
    print(f"[*] Text Overrides: {len(cfg['additional_adjustments']['text_overrides'])} aturan ditemukan")
    print("=========================================================")
