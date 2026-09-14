import openpyxl

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jul\Template_Pekerjaan_Xara_Jul.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Header & Ringkasan sheet
ws_h = wb['Header & Ringkasan']
print("=" * 80)
print("HEADER & RINGKASAN DATA:")
print("=" * 80)
for r in range(1, ws_h.max_row + 1):
    k = ws_h.cell(r, 1).value
    v = ws_h.cell(r, 2).value
    if k and v:
        print(f"  {k:30s} : {v}")

# 2. Tabel_Mutasi sheet
ws_m = wb['Tabel_Mutasi']
print("\n" + "=" * 80)
print("TABEL MUTASI AUDIT:")
print("=" * 80)

saldo_awal = 21347.81
dana_masuk = 0.0
dana_keluar = 0.0
running_bal = saldo_awal
valid_rows = []

for r in range(5, ws_m.max_row + 1):
    no = ws_m.cell(r, 1).value
    tgl = ws_m.cell(r, 2).value
    jam = ws_m.cell(r, 3).value
    uraian = ws_m.cell(r, 4).value
    nom = ws_m.cell(r, 5).value
    saldo_excel = ws_m.cell(r, 7).value
    
    if no == '[AWAL]':
        continue
        
    if nom is not None:
        nom_val = float(nom)
        if nom_val > 0:
            dana_masuk += nom_val
        else:
            dana_keluar += abs(nom_val)
        running_bal += nom_val
        valid_rows.append((no, tgl, jam, uraian, nom_val, running_bal, saldo_excel))

saldo_akhir = running_bal

print(f"Total Transactions in Excel: {len(valid_rows)}")
print(f"Saldo Awal   : Rp {saldo_awal:15,.2f}")
print(f"Dana Masuk   : Rp {dana_masuk:15,.2f}")
print(f"Dana Keluar  : Rp {dana_keluar:15,.2f}")
print(f"Saldo Akhir  : Rp {saldo_akhir:15,.2f}")
print(f"Balance Equation: Saldo Awal + Masuk - Keluar = {saldo_awal + dana_masuk - dana_keluar:,.2f} == Saldo Akhir {saldo_akhir:,.2f} -> {abs((saldo_awal + dana_masuk - dana_keluar) - saldo_akhir) < 0.01}")

print("\n--- FIRST 5 ROWS ---")
for r in valid_rows[:5]:
    print(r)

print("\n--- LAST 5 ROWS ---")
for r in valid_rows[-5:]:
    print(r)
