import openpyxl

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jul\Template_Pekerjaan_Xara_Jul.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n================================================================================")
    print(f"SHEET: {sname} (max_row={ws.max_row}, max_col={ws.max_column})")
    print(f"================================================================================")
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {row_vals}")
