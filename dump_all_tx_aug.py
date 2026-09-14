import openpyxl

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\aug\Template_Pekerjaan_Xara_Aug.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws_tx = wb['Tabel_Mutasi']

print(f"=== ALL NON-EMPTY ROWS IN Tabel_Mutasi (max_row={ws_tx.max_row}) ===")
for r in range(1, ws_tx.max_row + 1):
    vals = [ws_tx.cell(r, c).value for c in range(1, 10)]
    if any(v is not None for v in vals):
        print(f"Row {r:3d}: {vals}")
