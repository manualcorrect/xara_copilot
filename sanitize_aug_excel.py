import openpyxl

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\aug\Template_Pekerjaan_Xara_Aug.xlsx'
wb = openpyxl.load_workbook(excel_path)

ws_mutasi = wb['Tabel_Mutasi']

# Clean up corrupted cells in column B & C (rows 6..24)
# Restore clean dates
dates_fix = {
    6: ('01/08/2026', None),
    7: (None, None),
    8: (None, None),
    9: ('02/08/2026', None),
    10: (None, None),
    11: (None, None),
    12: (None, None),
    13: (None, None),
    14: ('05/08/2026', None),
    15: ('06/08/2026', None),
    16: ('08/08/2026', None),
    17: ('09/08/2026', None),
    18: (None, None),
    19: (None, None),
    20: (None, None),
    21: (None, None),
    22: (None, None),
    23: (None, None),
    24: (None, None),
    101: ('25/08/2026', '04:00:00 WIB'),
    115: ('31/08/2026', '23:59:00 WIB')
}

for r, (tgl, jam) in dates_fix.items():
    ws_mutasi.cell(r, 2).value = tgl
    ws_mutasi.cell(r, 3).value = jam

# Ensure views are not grouped
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = False
wb['Tabel_Mutasi'].sheet_view.tabSelected = True
wb.active = wb['Tabel_Mutasi']

wb.save(excel_path)
print("[*] Successfully sanitized Template_Pekerjaan_Xara_Aug.xlsx (Ungrouped & Cleaned Dates)!")
