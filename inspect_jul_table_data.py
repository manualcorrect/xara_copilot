import openpyxl

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jul\Template_Pekerjaan_Xara_Jul.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Tabel_Mutasi']

print("=" * 100)
print(f"INSPECTING TABEL_MUTASI (Total Rows: {ws.max_row})")
print("=" * 100)

for r in range(4, ws.max_row + 1):
    no = ws.cell(r, 1).value
    tgl = ws.cell(r, 2).value
    jam = ws.cell(r, 3).value
    uraian = ws.cell(r, 4).value
    nom = ws.cell(r, 5).value
    tipe = ws.cell(r, 6).value
    saldo = ws.cell(r, 7).value
    if any(v is not None for v in [no, tgl, uraian, nom, saldo]):
        print(f"Excel Row {r:2d} | No={str(no):6s} | Tgl={str(tgl):35s} | Jam={str(jam):12s} | Nom={str(nom):12s} | Saldo={str(saldo):14s} | Uraian={str(uraian)[:30]}")
