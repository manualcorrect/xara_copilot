import openpyxl
import json
import random
from datetime import datetime, time, timedelta

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jul\Template_Pekerjaan_Xara_Jul.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Tabel_Mutasi']

with open('jul_mapping.json', 'r', encoding='utf-8') as f:
    mapping = json.load(f)

# Extract raw transactions from Excel
excel_txs = []
current_date_str = "01 Jul 2026"

for r in range(5, ws.max_row + 1):
    no = ws.cell(r, 1).value
    tgl = ws.cell(r, 2).value
    jam = ws.cell(r, 3).value
    uraian = ws.cell(r, 4).value
    nom = ws.cell(r, 5).value
    saldo = ws.cell(r, 7).value
    
    if no == '[AWAL]':
        continue
    if nom is not None:
        # Determine Date
        if tgl is not None:
            tgl_s = str(tgl).strip()
            # If tgl_s is corrupted header string (like path or label), ignore and keep current_date_str
            if '/' in tgl_s:
                parts = tgl_s.split('/')
                day = int(parts[0])
                month = int(parts[1])
                current_date_str = f"{day:02d} Jul 2026"
            elif isinstance(tgl, datetime):
                current_date_str = f"{tgl.day:02d} Jul 2026"
            elif '2026-' in tgl_s:
                dt = datetime.strptime(tgl_s.split()[0], '%Y-%m-%d')
                current_date_str = f"{dt.day:02d} Jul 2026"
                
        # Determine Time
        time_str = None
        if jam is not None:
            jam_s = str(jam).strip()
            if 'WIB' in jam_s:
                time_str = jam_s
            elif isinstance(jam, time):
                time_str = f"{jam.strftime('%H:%M:%S')} WIB"
                
        excel_txs.append({
            "index": len(excel_txs) + 1,
            "excel_row": r,
            "date_str": current_date_str,
            "explicit_time": time_str,
            "nominal": float(nom),
            "saldo": float(saldo)
        })

print(f"Parsed {len(excel_txs)} transactions from Excel")

# Group by date to generate ascending smart timestamps
random.seed(42) # Deterministic realistic seed
date_groups = {}
for tx in excel_txs:
    d = tx["date_str"]
    if d not in date_groups:
        date_groups[d] = []
    date_groups[d].append(tx)

for d, tx_list in date_groups.items():
    n_tx = len(tx_list)
    start_sec = 6 * 3600 + 15 * 60 # 06:15:00
    end_sec = 22 * 3600 + 45 * 60   # 22:45:00
    
    # Generate n_tx strictly ascending random timestamps
    raw_secs = sorted(random.sample(range(start_sec, end_sec), n_tx))
    for i, tx in enumerate(tx_list):
        if tx["explicit_time"]:
            tx["final_time"] = tx["explicit_time"]
        else:
            sec = raw_secs[i]
            hh = sec // 3600
            mm = (sec % 3600) // 60
            ss = sec % 60
            tx["final_time"] = f"{hh:02d}:{mm:02d}:{ss:02d} WIB"

# Verify 83 rows mapped to mapping["rows"]
assert len(excel_txs) == len(mapping["rows"]), f"Mismatch! Excel {len(excel_txs)} vs XAR {len(mapping['rows'])}"

for i in range(len(excel_txs)):
    excel_txs[i]["xar_row_info"] = mapping["rows"][i]

with open('jul_final_tx_schedule.json', 'w', encoding='utf-8') as f:
    json.dump(excel_txs, f, indent=2)

print("\n--- FIRST 5 SCHEDULED ROWS ---")
for tx in excel_txs[:5]:
    print(f"Row {tx['index']:2d} | Date: {tx['date_str']} | Time: {tx['final_time']} | Nom: {tx['nominal']:12,.2f} | Saldo: {tx['saldo']:14,.2f}")

print("\n--- LAST 5 SCHEDULED ROWS ---")
for tx in excel_txs[-5:]:
    print(f"Row {tx['index']:2d} | Date: {tx['date_str']} | Time: {tx['final_time']} | Nom: {tx['nominal']:12,.2f} | Saldo: {tx['saldo']:14,.2f}")
