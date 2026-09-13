import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import shutil

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styles Definition
font_family = "Segoe UI"
header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Deep Navy
section_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Blue Accent
zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
green_font = Font(name=font_family, size=10, color="276A3C", bold=True)
red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
red_font = Font(name=font_family, size=10, color="A61C1C", bold=True)
blue_font = Font(name=font_family, size=10, color="005B9C", bold=True)
gold_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Soft Ice Blue Accent for Saldo Awal
gold_font = Font(name=font_family, size=10, color="1B365D", bold=True)
total_fill = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")

title_font = Font(name=font_family, size=14, bold=True, color="1B365D")
subtitle_font = Font(name=font_family, size=9, italic=True, color="595959")
th_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
section_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
bold_font = Font(name=font_family, size=10, bold=True)
regular_font = Font(name=font_family, size=10)
italic_font = Font(name=font_family, size=9, italic=True, color="7F7F7F")

thin_border_side = Side(style='thin', color="D9D9D9")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
thick_top_bottom = Border(
    top=Side(style='medium', color="1B365D"),
    bottom=Side(style='double', color="1B365D"),
    left=thin_border_side,
    right=thin_border_side
)

align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# =========================================================================
# SHEET 1: Header & Ringkasan
# =========================================================================
ws1 = wb.create_sheet(title="Header & Ringkasan")
ws1.views.sheetView[0].showGridLines = True

ws1["A1"] = "KONFIGURASI UTAMA PEKERJAAN (TAHAP 1 - 5 & RINGKASAN TAHAP 7)"
ws1["A1"].font = title_font
ws1["A2"] = "Template ini diisi sesuai data pekerjaan Anda. Data di bawah ini memuat konfigurasi Stress Test 3 (test_3.1.xar)."
ws1["A2"].font = subtitle_font

headers1 = ["Kategori / Parameter", "Nilai Data (Input Anda)", "Status", "Keterangan & Aturan Kerja (SOP Project V2)"]
for col_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=4, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = th_font
    cell.alignment = align_center

params_s1 = [
    ("SEC", "1. INFORMASI FILE & PATH PROYEK", "", ""),
    ("File Sumber (.xar)", r"C:\Users\Lenovo\Downloads\rekening\Antigravity\test\Test_2\test_3.1.xar", "Wajib", "Lokasi file .xar template asli yang akan diedit"),
    ("File Output (.xar)", r"C:\Users\Lenovo\Downloads\rekening\Antigravity\test\Test_2\test_3.1_output.xar", "Wajib", "Lokasi file .xar baru hasil editan (file asli tetap aman)"),
    ("Nama Dokumen / Catatan", "Rekening Koran Desember 2026 - ASEP ISKANDAR", "Opsional", "Keterangan label proyek untuk arsip riwayat training"),
    
    ("SEC", "2. DATA HEADER NASABAH (TAHAP 1 s.d. 4)", "", ""),
    ("Nama Nasabah (Tahap 1)", "ASEP ISKANDAR", "Wajib", "Mengubah Nama/Name di semua lembar. Otomatis diberi spasi dan newline '\\r\\n' agar cabang tetap di baris 2"),
    ("Periode Laporan (Tahap 2)", "01 Dec 2026 - 31 Dec 2026", "Wajib", "Mengubah rentang periode di seluruh halaman header secara otomatis"),
    ("Dicetak Pada (Tahap 3)", "19 Jan 2027", "Wajib", "Mengubah tanggal cetak / Issued on di seluruh lembar halaman"),
    ("Nomor Rekening (Tahap 4)", "1630000000000", "Wajib", "Nomor rekening nasabah (13 digit atau sesuai buku rekening)"),
    
    ("SEC", "3. PENOMORAN HALAMAN (TAHAP 5)", "", ""),
    ("Mode Total Halaman", "AUTO", "Direkomendasikan", "Ketik 'AUTO' agar sistem mendeteksi total lembar secara mandiri (misal 5 lembar -> 1 dari 5 s.d. 5 dari 5)"),
    ("Total Halaman Manual", "5", "Opsional", "Hanya digunakan jika ingin memaksa angka total halaman tertentu"),
    
    ("SEC", "4. BULAN & TAHUN TRANSAKSI (TAHAP 6)", "", ""),
    ("Target Bulan & Tahun", "Dec 2026", "Wajib", "Mengganti bulan dan tahun seluruh baris mutasi tabel secara serempak"),
    ("Digit Penutup Tahun", "6", "Wajib", "Mengganti node digit penutup tahun (misal tahun 2026 -> digit 6)"),
    
    ("SEC", "5. RINGKASAN KEUANGAN HEADER (TAHAP 7)", "", ""),
    ("Saldo Awal", "26.683,00", "Wajib", "Saldo awal rekening awal periode. Nilai ini menjadi saldo pembuka baris pertama pada Sheet 'Tabel_Mutasi'"),
    ("Dana Masuk (Kredit)", "+ 5.315.920,00", "Wajib", "Total mutasi uang masuk pada periode ini (Warna Hijau #00A651)"),
    ("Dana Keluar (Debit)", "- 4.210.600,00", "Wajib", "Total mutasi uang keluar pada periode ini (Warna Hitam #000000)"),
    ("Saldo Akhir", "1.132.003,00", "Wajib", "Saldo penutupan akhir periode (Warna Biru #005B9C)"),
    ("Keseimbangan Neraca (Audit)", "BALANCE (MATCH)", "Auto-Check", "Formula audit: Saldo Awal (26.683) + Masuk (5.315.920) - Keluar (4.210.600) = 1.132.003,00")
]

cur_r = 5
for row_data in params_s1:
    if row_data[0] == "SEC":
        ws1.merge_cells(start_row=cur_r, start_column=1, end_row=cur_r, end_column=4)
        c = ws1.cell(row=cur_r, column=1, value=row_data[1])
        c.fill = section_fill
        c.font = section_font
        c.alignment = align_left
        for col in range(1, 5):
            ws1.cell(row=cur_r, column=col).border = thin_border
    else:
        c1 = ws1.cell(row=cur_r, column=1, value=row_data[0])
        c2 = ws1.cell(row=cur_r, column=2, value=row_data[1])
        c3 = ws1.cell(row=cur_r, column=3, value=row_data[2])
        c4 = ws1.cell(row=cur_r, column=4, value=row_data[3])
        
        c1.font = bold_font
        c2.font = bold_font if "Wajib" in row_data[2] else regular_font
        c3.font = italic_font
        c4.font = regular_font
        
        c1.border = thin_border
        c2.border = thin_border
        c3.border = thin_border
        c4.border = thin_border
        
        c1.alignment = align_left
        c2.alignment = align_left
        c3.alignment = align_center
        c4.alignment = align_left
        
        # Color coding for Ringkasan values
        if row_data[0] == "Dana Masuk (Kredit)":
            c2.font = green_font
            c2.fill = green_fill
        elif row_data[0] == "Dana Keluar (Debit)":
            c2.font = red_font
            c2.fill = red_fill
        elif row_data[0] == "Saldo Akhir":
            c2.font = blue_font
        elif row_data[0] == "Keseimbangan Neraca (Audit)":
            c2.font = green_font
            c2.fill = green_fill
            c3.font = bold_font
    cur_r += 1


# =========================================================================
# SHEET 2: Tabel_Mutasi
# =========================================================================
ws2 = wb.create_sheet(title="Tabel_Mutasi")
ws2.views.sheetView[0].showGridLines = True

ws2["A1"] = "TABEL MUTASI TRANSAKSI UTAMA (47 BARIS TRANSAKSI)"
ws2["A1"].font = title_font
ws2["A2"] = "Daftar mutasi rekening koran. Baris pertama memuat Saldo Awal pembukaan, diikuti rincian transaksi 1 s.d. 47."
ws2["A2"].font = subtitle_font

headers2 = [
    "No", "Tanggal", "Jam", "Uraian Transaksi", "Nominal (+/-)", 
    "Tipe (CR/DB)", "Saldo Berjalan", "Kustomisasi Warna", "Kustomisasi Alignment", "Catatan Baris"
]

for col_idx, h in enumerate(headers2, 1):
    cell = ws2.cell(row=4, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = th_font
    cell.alignment = align_center

# ROW 5: Dedicated Saldo Awal Row (Memperjelas hubungan Saldo Awal dengan Sheet Mutasi)
row_awal = [
    "[AWAL]", "01/12", "-", "SALDO AWAL PEMBUKAAN PERIODE (INITIAL BALANCE)", "-", 
    "AWAL", "26.683,00", "AUTO (Tag 150)", "AUTO (20.049 cm)", "Saldo awal pembukaan rekening periode laporan"
]

for c_idx, val in enumerate(row_awal, 1):
    cell = ws2.cell(row=5, column=c_idx, value=val)
    cell.fill = gold_fill
    cell.font = gold_font
    cell.border = thin_border
    if c_idx in (1, 2, 3, 6, 8, 9):
        cell.alignment = align_center
    elif c_idx in (5, 7):
        cell.alignment = align_right
    else:
        cell.alignment = align_left

user_table_47 = [
    (1, '01/12', '04:00:00 WIB', 'TRSF E-BANKING CR / PENGEMBALIAN DANA', '+15.000,00', 'CR', '41.683,00'),
    (2, '01/12', '04:12:15 WIB', 'TRSF E-BANKING DB / BIAYA TRANSFER', '-16.000,00', 'DB', '25.683,00'),
    (3, '02/12', '08:30:00 WIB', 'SETORAN TUNAI VIA CRM', '+700.000,00', 'CR', '725.683,00'),
    (4, '02/12', '09:15:22 WIB', 'BIAYA ADM REKENING KORAN', '-2.500,00', 'DB', '723.183,00'),
    (5, '03/12', '10:05:10 WIB', 'TARIK TUNAI ATM LINK', '-500.000,00', 'DB', '223.183,00'),
    (6, '04/12', '11:20:45 WIB', 'TRANSFER SESAMA BANK MANDIRI', '-100.000,00', 'DB', '123.183,00'),
    (7, '04/12', '12:00:00 WIB', 'BIAYA TRANSAKSI BI-FAST', '-2.000,00', 'DB', '121.183,00'),
    (8, '05/12', '13:45:12 WIB', 'PEMBELIAN PULSA PRABAYAR', '-15.000,00', 'DB', '106.183,00'),
    (9, '05/12', '14:10:00 WIB', 'PEMBAYARAN TAGIHAN LISTRIK PLN', '-11.000,00', 'DB', '95.183,00'),
    (10, '06/12', '15:25:30 WIB', 'BIAYA ADMIN BULANAN', '-2.500,00', 'DB', '92.683,00'),
    (11, '07/12', '07:10:05 WIB', 'QRIS MERCHANT INDOMARET', '-15.000,00', 'DB', '77.683,00'),
    (12, '08/12', '08:50:11 WIB', 'BIAYA NOTIFIKASI SMS', '-2.000,00', 'DB', '75.683,00'),
    (13, '08/12', '09:30:00 WIB', 'QRIS MERCHANT KOPI KENANGAN', '-15.000,00', 'DB', '60.683,00'),
    (14, '09/12', '10:00:00 WIB', 'TRANSFER MASUK DARI BCA', '+300.000,00', 'CR', '360.683,00'),
    (15, '10/12', '11:15:20 WIB', 'TRANSFER KE SHOPEEPAY TOPUP', '-300.000,00', 'DB', '60.683,00'),
    (16, '11/12', '12:35:40 WIB', 'PEMBELIAN PAKET DATA INTERNET', '-19.000,00', 'DB', '41.683,00'),
    (17, '12/12', '13:10:00 WIB', 'CASHBACK TRANSAKSI BULANAN', '+214.000,00', 'CR', '255.683,00'),
    (18, '12/12', '14:20:15 WIB', 'TARIK TUNAI CRM', '-30.000,00', 'DB', '225.683,00'),
    (19, '13/12', '15:05:00 WIB', 'PEMBAYARAN ASURANSI KESEHATAN', '-200.000,00', 'DB', '25.683,00'),
    (20, '14/12', '09:00:00 WIB', 'SETORAN TUNAI BANK MANDIRI', '+250.000,00', 'CR', '275.683,00'),
    (21, '15/12', '10:45:10 WIB', 'TRANSFER KE REKENING TABUNGAN', '-200.000,00', 'DB', '75.683,00'),
    (22, '15/12', '11:12:00 WIB', 'QRIS RESTORAN PADANG', '-18.100,00', 'DB', '57.583,00'),
    (23, '16/12', '12:00:15 WIB', 'PARKIR NON TUNAI OVO', '-6.800,00', 'DB', '50.783,00'),
    (24, '17/12', '13:30:20 WIB', 'BIAYA TRANSAKSI BI-FAST', '-2.000,00', 'DB', '48.783,00'),
    (25, '18/12', '14:15:00 WIB', 'PEMBELIAN PULSA INDOSAT', '-15.000,00', 'DB', '33.783,00'),
    (26, '19/12', '15:40:10 WIB', 'QRIS MINIMARKET ALFAMART', '-5.700,00', 'DB', '28.083,00'),
    (27, '20/12', '08:30:00 WIB', 'DANA MASUK GAJI BONUS TAHUNAN', '+300.000,00', 'CR', '328.083,00'),
    (28, '21/12', '09:20:00 WIB', 'PEMBAYARAN KARTU KREDIT', '-300.000,00', 'DB', '28.083,00'),
    (29, '22/12', '10:15:00 WIB', 'TRANSFER MASUK HASIL USAHA CV', '+3.235.920,00', 'CR', '3.264.003,00'),
    (30, '22/12', '11:45:00 WIB', 'PEMBAYARAN TAGIHAN PDAM', '-70.000,00', 'DB', '3.194.003,00'),
    (31, '23/12', '12:10:00 WIB', 'BIAYA CEK SALDO ATM BERSAMA', '-500,00', 'DB', '3.193.503,00'),
    (32, '23/12', '13:00:00 WIB', 'PEMBAYARAN BPJS KESEHATAN', '-97.000,00', 'DB', '3.096.503,00'),
    (33, '24/12', '14:20:00 WIB', 'BIAYA CEK SALDO ATM BERSAMA', '-500,00', 'DB', '3.096.003,00'),
    (34, '24/12', '15:10:00 WIB', 'QRIS COFFEE SHOP POINT', '-15.000,00', 'DB', '3.081.003,00'),
    (35, '25/12', '09:00:00 WIB', 'TARIK TUNAI ATM MANDIRI', '-100.000,00', 'DB', '2.981.003,00'),
    (36, '26/12', '10:30:00 WIB', 'DANA MASUK TALANGAN PROYEK', '+300.000,00', 'CR', '3.281.003,00'),
    (37, '26/12', '11:00:00 WIB', 'BUNGA TABUNGAN BANK', '+1.000,00', 'CR', '3.282.003,00'),
    (38, '27/12', '12:15:00 WIB', 'PAJAK BUNGA TABUNGAN', '-36.500,00', 'DB', '3.245.503,00'),
    (39, '27/12', '13:45:00 WIB', 'TRANSFER KE REK KELUARGA', '-200.000,00', 'DB', '3.045.503,00'),
    (40, '28/12', '14:10:00 WIB', 'PEMBAYARAN CICILAN KREDIT', '-100.000,00', 'DB', '2.945.503,00'),
    (41, '28/12', '15:00:00 WIB', 'BELANJA MARKETPLACE TOKOPEDIA', '-100.000,00', 'DB', '2.845.503,00'),
    (42, '29/12', '08:45:00 WIB', 'TRANSFER KE INVESTASI REKSADANA', '-500.000,00', 'DB', '2.345.503,00'),
    (43, '29/12', '09:30:00 WIB', 'BIAYA MATERAI ELEKTRONIK', '-3.500,00', 'DB', '2.342.003,00'),
    (44, '30/12', '10:15:00 WIB', 'TARIK TUNAI REKENING KORAN', '-200.000,00', 'DB', '2.142.003,00'),
    (45, '30/12', '11:40:00 WIB', 'PEMBAYARAN SEWA KANTOR CABANG', '-1.000.000,00', 'DB', '1.142.003,00'),
    (46, '31/12', '13:00:00 WIB', 'BIAYA PEMELIHARAAN KARTU DEBIT', '-5.000,00', 'DB', '1.137.003,00'),
    (47, '31/12', '14:30:00 WIB', 'BIAYA ADM LAYANAN INTERNET BANKING', '-5.000,00', 'DB', '1.132.003,00')
]

for idx, rdata in enumerate(user_table_47, 1):
    row_idx = 5 + idx
    num, tgl, jam, uraian, nom, tipe, saldo = rdata
    
    c_num = ws2.cell(row=row_idx, column=1, value=num)
    c_tgl = ws2.cell(row=row_idx, column=2, value=tgl)
    c_jam = ws2.cell(row=row_idx, column=3, value=jam)
    c_ura = ws2.cell(row=row_idx, column=4, value=uraian)
    c_nom = ws2.cell(row=row_idx, column=5, value=nom)
    c_tip = ws2.cell(row=row_idx, column=6, value=tipe)
    c_sal = ws2.cell(row=row_idx, column=7, value=saldo)
    c_col = ws2.cell(row=row_idx, column=8, value="AUTO (Tag 150)")
    c_ali = ws2.cell(row=row_idx, column=9, value="AUTO (15.214 cm)")
    c_cat = ws2.cell(row=row_idx, column=10, value="OK Verified")
    
    # Alignments
    c_num.alignment = align_center
    c_tgl.alignment = align_center
    c_jam.alignment = align_center
    c_ura.alignment = align_left
    c_nom.alignment = align_right
    c_tip.alignment = align_center
    c_sal.alignment = align_right
    c_col.alignment = align_center
    c_ali.alignment = align_center
    c_cat.alignment = align_left
    
    # Borders
    for col in range(1, 11):
        cell_curr = ws2.cell(row=row_idx, column=col)
        cell_curr.border = thin_border
        cell_curr.font = regular_font
        if row_idx % 2 == 1:
            cell_curr.fill = zebra_fill
        
    c_num.font = bold_font
    c_sal.font = blue_font
    
    if tipe == 'CR':
        c_nom.font = green_font
        c_tip.font = green_font
        c_nom.fill = green_fill
    else:
        c_nom.font = red_font
        c_tip.font = red_font
        c_nom.fill = red_fill

# ROW 53: Summary Total Mutasi & Saldo Akhir Row
row_total_idx = 5 + len(user_table_47) + 1
row_total_data = [
    "[TOTAL]", "31/12", "-", "TOTAL MUTASI & SALDO PENUTUPAN AKHIR", "+ 1.105.320,00",
    "NET", "1.132.003,00", "AUTO (Tag 150)", "AUTO (20.049 cm)", "Audit: Masuk +5.315.920,00 | Keluar -4.210.600,00 | MATCH"
]

for c_idx, val in enumerate(row_total_data, 1):
    cell = ws2.cell(row=row_total_idx, column=c_idx, value=val)
    cell.fill = total_fill
    cell.font = bold_font
    cell.border = thick_top_bottom
    if c_idx in (1, 2, 3, 6, 8, 9):
        cell.alignment = align_center
    elif c_idx in (5, 7):
        cell.alignment = align_right
    else:
        cell.alignment = align_left

ws2.cell(row=row_total_idx, column=7).font = blue_font


# =========================================================================
# SHEET 3: Penyesuaian_Tambahan
# =========================================================================
ws3 = wb.create_sheet(title="Penyesuaian_Tambahan")
ws3.views.sheetView[0].showGridLines = True

ws3["A1"] = "LEMBAR KHUSUS PENYESUAIAN TAMBAHAN (ADVANCED CUSTOMIZATION)"
ws3["A1"].font = title_font
ws3["A2"] = "Sheet ini khusus digunakan jika dokumen Anda memerlukan modifikasi tingkat lanjut (ruler khusus, warna kustom, override teks, dll)."
ws3["A2"].font = subtitle_font

headers3 = ["Kategori Penyesuaian", "Nama Parameter / Teks Asal", "Nilai Kustom / Teks Baru", "Target Halaman / Elemen", "Keterangan & Panduan"]
for col_idx, h in enumerate(headers3, 1):
    cell = ws3.cell(row=4, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = th_font
    cell.alignment = align_center

params_s3 = [
    ("SEC", "A. KUSTOMISASI RULER & BATAS RATA KANAN (RIGHT-ALIGNMENT GRID)", "", "", ""),
    ("Ruler Grid", "Target X-Right Kolom Nominal", "431320 mp (15.214 cm)", "Semua Baris Nominal", "Titik koordinat batas sisi kanan kolom Nominal pada Xara"),
    ("Ruler Grid", "Target X-Right Kolom Saldo", "570450 mp (20.049 cm)", "Semua Baris Saldo", "Titik koordinat batas sisi kanan kolom Saldo pada Xara"),
    ("Ruler Grid", "Mode Kalkulasi Otomatis Width", "AKTIF", "Tag 2206 & 2100", "Menghitung mundur X_left = X_target - W(teks) sesuai glyph metrics font"),
    
    ("SEC", "B. KUSTOMISASI PALET WARNA BINER (TAG 150)", "", "", ""),
    ("Warna Biner", "Nominal Kredit (CR)", "b'\\x0b\\x04\\x00\\x00' (#00A651)", "Tabel Mutasi", "Kunci warna Hijau resmi bank"),
    ("Warna Biner", "Nominal Debit (DB)", "b'\\x95\\x0e\\x00\\x00' (#000000)", "Tabel Mutasi", "Kunci warna Hitam resmi bank"),
    ("Warna Biner", "Saldo Berjalan", "b'\\x88\\x05\\x00\\x00' (#005B9C)", "Tabel Mutasi", "Kunci warna Biru resmi bank"),
    ("Warna Biner", "Header Dana Masuk", "b'\\x0b\\x04\\x00\\x00' (#00A651)", "Header Ringkasan", "Kunci warna Hijau ringkasan"),
    ("Warna Biner", "Header Dana Keluar", "b'\\x76\\x02\\x00\\x00' (#000000)", "Header Ringkasan", "Kunci warna Hitam ringkasan"),
    ("Warna Biner", "Header Saldo Akhir", "b'\\x88\\x05\\x00\\x00' (#005B9C)", "Header Ringkasan", "Kunci warna Biru ringkasan"),
    
    ("SEC", "C. OVERRIDE TEKS KHUSUS / ALAMAT / CABANG TAMBAHAN", "", "", ""),
    ("Override Teks", "KCP JAKARTA SUDIRMAN", "KCP JAKARTA SUDIRMAN", "Header Lembar 1", "Dapat diubah jika ingin mengganti nama kantor cabang"),
    ("Override Teks", "JL. JEND. SUDIRMAN KAV. 52-53", "JL. JEND. SUDIRMAN KAV. 52-53", "Header Lembar 1", "Dapat diubah jika ingin mengganti alamat nasabah"),
    ("Override Teks", "JAKARTA SELATAN 12190", "JAKARTA SELATAN 12190", "Header Lembar 1", "Dapat diubah jika ingin mengganti kota dan kode pos"),
    ("Override Teks", "IDR", "IDR", "Semua Halaman", "Mata uang laporan rekening"),
    ("Override Teks", "[Teks Kustom Lain]", "[Pengganti Kustom] ", "[Halaman Target]", "Slot bebas untuk teks tambahan lain yang ingin diganti"),
    
    ("SEC", "D. FITUR KEAMANAN BINER & ANTI-CORRUPTION (STABILITY SWITCHES)", "", "", ""),
    ("Keamanan Biner", "Zero-Shift Pointer Mandate", "TERKUNCI (AKTIF)", "Seluruh Stream", "Mencegah penambahan/pengurangan record agar ID warna & font tidak rusak"),
    ("Keamanan Biner", "Tag 2202 Null Character Safety", "b'\\x00\\x00' (2 Bytes)", "Pecahan Node", "Menjamin node sekunder tidak pernah 0 byte agar Xara tidak crash"),
    ("Keamanan Biner", "Auto-Sync Payload Size", "AKTIF", "Header Tiap Record", "rec['size'] = len(rec['payload']) selalu sinkron sebelum file disimpan"),
    ("Keamanan Biner", "Lock Embedded Font (Tag 2907)", "TERKUNCI (AKTIF)", "Font Definitions", "Mencegah font fallback ter-reset menjadi default PDF"),
    
    ("SEC", "E. LOG CATATAN OPERATOR / INSTRUKSI KHUSUS PEKERJAAN", "", "", ""),
    ("Catatan Operator", "Catatan Klien / Khusus", "Tidak ada instruksi khusus", "Audit Trail", "Catatan bebas untuk operator sebelum dokumen diekspor")
]

cur_r3 = 5
for row_data in params_s3:
    if row_data[0] == "SEC":
        ws3.merge_cells(start_row=cur_r3, start_column=1, end_row=cur_r3, end_column=5)
        c = ws3.cell(row=cur_r3, column=1, value=row_data[1])
        c.fill = section_fill
        c.font = section_font
        c.alignment = align_left
        for col in range(1, 6):
            ws3.cell(row=cur_r3, column=col).border = thin_border
    else:
        c1 = ws3.cell(row=cur_r3, column=1, value=row_data[0])
        c2 = ws3.cell(row=cur_r3, column=2, value=row_data[1])
        c3 = ws3.cell(row=cur_r3, column=3, value=row_data[2])
        c4 = ws3.cell(row=cur_r3, column=4, value=row_data[3])
        c5 = ws3.cell(row=cur_r3, column=5, value=row_data[4])
        
        c1.font = bold_font
        c2.font = regular_font
        c3.font = bold_font if "AKTIF" in row_data[2] or "TERKUNCI" in row_data[2] else regular_font
        c4.font = italic_font
        c5.font = regular_font
        
        for col in range(1, 6):
            ws3.cell(row=cur_r3, column=col).border = thin_border
            
        c1.alignment = align_left
        c2.alignment = align_left
        c3.alignment = align_left
        c4.alignment = align_center
        c5.alignment = align_left
    cur_r3 += 1

# Column Widths Setup
ws1.column_dimensions['A'].width = 30
ws1.column_dimensions['B'].width = 45
ws1.column_dimensions['C'].width = 18
ws1.column_dimensions['D'].width = 65

ws2.column_dimensions['A'].width = 14
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 16
ws2.column_dimensions['D'].width = 45
ws2.column_dimensions['E'].width = 18
ws2.column_dimensions['F'].width = 14
ws2.column_dimensions['G'].width = 18
ws2.column_dimensions['H'].width = 20
ws2.column_dimensions['I'].width = 22
ws2.column_dimensions['J'].width = 22

ws3.column_dimensions['A'].width = 22
ws3.column_dimensions['B'].width = 35
ws3.column_dimensions['C'].width = 35
ws3.column_dimensions['D'].width = 25
ws3.column_dimensions['E'].width = 50

# Save to both target locations
out_path_1 = r"C:\Users\Lenovo\xara_copilot\Template_Pekerjaan_Xara.xlsx"
out_path_2 = r"C:\Users\Lenovo\Downloads\rekening\Antigravity\test\Test_2\Template_Pekerjaan_Xara.xlsx"

wb.save(out_path_1)
print(f"Successfully saved template at: {out_path_1}")

try:
    shutil.copy(out_path_1, out_path_2)
    print(f"Successfully copied template to: {out_path_2}")
except PermissionError:
    print(f"Catatan: {out_path_2} sedang dibuka di Microsoft Excel. Silakan tutup file tersebut jika ingin menyegarkan salinannya.")
