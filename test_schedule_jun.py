import openpyxl
from datetime import datetime
from smart_date_time_generator import generate_smart_schedule

wb = openpyxl.load_workbook(r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jun\Template_Pekerjaan_Xara_Jun.xlsx', data_only=True)
ws = wb['Tabel_Mutasi']

anchors = {}
for r in range(6, 79):
    idx = r - 6
    raw_d = ws.cell(r, 2).value
    raw_t = ws.cell(r, 3).value
    
    entry = {}
    if raw_d:
        if isinstance(raw_d, datetime):
            if raw_d.month == 6:
                d_num = min(30, max(1, raw_d.day))
                entry['date'] = f"2026-06-{d_num:02d}"
        elif '/' in str(raw_d):
            p = str(raw_d).strip().split('/')
            d_num = min(30, max(1, int(p[0])))
            entry['date'] = f"2026-06-{d_num:02d}"
    if raw_t:
        t_str = str(raw_t).strip()
        if 'WIB' not in t_str:
            t_str += ' WIB'
        entry['time'] = t_str
    if entry:
        anchors[idx] = entry

print(f"Total anchors: {len(anchors)}")
for k, v in sorted(anchors.items()):
    print(f"  Row {k+1:2d}: {v}")

sched = generate_smart_schedule(73, period_start="2026-06-01", period_end="2026-06-30", manual_anchors=anchors)
print("\nGenerated Smart Schedule:")
for i in range(73):
    print(f"Row {i+1:2d}: {sched[i]['date_full']} | {sched[i]['time']}")

