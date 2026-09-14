import openpyxl
import os
import struct
from xar_dom_engine import XarDocument

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jul\Template_Pekerjaan_Xara_Jul.xlsx'
xar_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jul\0.xar'

print("=" * 80)
print("1. INSPECTING EXCEL TEMPLATE")
print("=" * 80)

wb = openpyxl.load_workbook(excel_path, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- Sheet: {sheetname} ({ws.max_row} rows, {ws.max_column} cols) ---")
    for r in range(1, min(ws.max_row + 1, 20)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 10))]
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {row_vals}")

print("\n" + "=" * 80)
print("2. INSPECTING 0.XAR DOM")
print("=" * 80)

doc = XarDocument(xar_path)
total_recs = len(doc.records)
print(f"Total records in 0.xar: {total_recs:,}")

# Count pages
page_tags = [i for i, r in enumerate(doc.records) if r['tag'] == 4465 or (r['tag'] == 2201 and 'Page ' in r['payload'].decode('utf-16le', errors='ignore'))]
print(f"Found {len(page_tags)} page markers/records")

# Search all text in 0.xar
texts = []
for i, r in enumerate(doc.records):
    if r['tag'] in (2201, 2202):
        t = r['payload'].decode('utf-16le', errors='ignore').rstrip('\x00')
        if t:
            texts.append((i, r['tag'], t))

print(f"Total non-empty text nodes: {len(texts)}")
for i, tag, t in texts[:35]:
    print(f"Rec {i:5d} [Tag {tag}]: '{t}'")
