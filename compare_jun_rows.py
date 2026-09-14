import json
import openpyxl

with open('jun_rows_map.json') as f:
    rows = json.load(f)

print("Mapped row numbers in 0.xar:", [r['row_no_val'] for r in rows])

wb = openpyxl.load_workbook(r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jun\Template_Pekerjaan_Xara_Jun.xlsx', data_only=True)
ws = wb['Tabel_Mutasi']
excel_rows = []
for r in range(5, ws.max_row+1):
    no = ws.cell(r, 1).value
    nom = ws.cell(r, 5).value
    saldo = ws.cell(r, 7).value
    if nom is not None and nom != '-':
        excel_rows.append((no, nom, saldo))

print(f"Total excel rows: {len(excel_rows)}")
print("Excel row numbers:", [r[0] for r in excel_rows])
