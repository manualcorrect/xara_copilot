import openpyxl
import os

excel_path = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\jun\Template_Pekerjaan_Xara_Jun.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("=== JUN EXCEL SHEETS ===", wb.sheetnames)
ws_hdr = wb['Header & Ringkasan']
print("Header & Ringkasan:")
for r in range(1, 26):
    row_vals = [ws_hdr.cell(r, c).value for c in range(1, 5)]
    if any(v is not None for v in row_vals):
        print(f"  Row {r:2d}: {row_vals}")

ws_mutasi = wb['Tabel_Mutasi']
tx_count = 0
for r in range(5, ws_mutasi.max_row+1):
    nom = ws_mutasi.cell(r, 5).value
    if nom is not None and nom != '-':
        tx_count += 1

print(f"\nTotal transactions in Tabel_Mutasi: {tx_count}")
