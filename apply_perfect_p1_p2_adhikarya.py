"""
PERFECT RUNNER FOR PAGE 1 & PAGE 2: ADHIKARYA PUTRA JUN 2026
Applies all 5 SOP/Knowledge Base rules + Page 1 & 2 Tahap 7 with flawless precision.
"""

import os
import sys
import json
import struct
import openpyxl
from xar_dom_engine import XarDocument

# Ensure unbuffered UTF-8 output
sys.stdout.reconfigure(encoding='utf-8', line_buffering=True, errors='replace')

# TTInterphases-Bold glyph widths
CHAR_WIDTHS = {
    '0': 5438, '1': 3160, '2': 4762, '3': 4840, '4': 5039,
    '5': 4840, '6': 4878, '7': 4402, '8': 4962, '9': 4878,
    '.': 1840, ',': 1243, '-': 3198, '+': 4399, ' ': 2200
}

def calc_text_width(text):
    return sum(CHAR_WIDTHS.get(c, 4800) for c in text)

TARGET_XR_NOMINAL = 431250 # 15.214 cm

# Native Palette
COLOR_GREEN = bytearray.fromhex('ee030000') # Hijau CR / Dana Masuk
COLOR_BLACK = bytearray.fromhex('9e010000') # Hitam DB / Dana Keluar
COLOR_BLUE  = bytearray.fromhex('50050000') # Biru Saldo
COLOR_GRAY  = bytearray.fromhex('8a030000') # Abu Saldo Awal

def fmt_idr(val):
    if isinstance(val, (int, float)):
        num = float(val)
    else:
        s = str(val).strip().replace('+', '').replace('-', '').replace(' ', '')
        if ',' in s and '.' in s:
            s = s.replace('.', '').replace(',', '.')
        elif ',' in s:
            s = s.replace(',', '.')
        num = float(s or 0.0)
    return f"{abs(num):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

def update_text(doc, rec_idx, text_str):
    p = bytearray(text_str.encode('utf-16le'))
    doc.records[rec_idx]['payload'] = p
    doc.records[rec_idx]['size'] = len(p)

def blank_node(doc, rec_idx):
    if rec_idx is not None:
        doc.records[rec_idx]['payload'] = bytearray(b'\x00\x00')
        doc.records[rec_idx]['size'] = 2

def update_color(doc, rec_idx, color_bytes):
    if rec_idx is not None:
        doc.records[rec_idx]['payload'] = color_bytes
        doc.records[rec_idx]['size'] = len(color_bytes)

def update_t2206(doc, rec_idx, new_w):
    if rec_idx is not None:
        orig = struct.unpack('<iii', doc.records[rec_idx]['payload'][:12])
        doc.records[rec_idx]['payload'] = bytearray(struct.pack('<iii', new_w, orig[1], orig[2]))
        doc.records[rec_idx]['size'] = 12

def update_t2206_full(doc, rec_idx, w, h, dx=0):
    if rec_idx is not None:
        doc.records[rec_idx]['payload'] = bytearray(struct.pack('<iii', w, h, dx))
        doc.records[rec_idx]['size'] = 12

def update_t2100_x(doc, rec_idx, new_x):
    if rec_idx is not None:
        orig = struct.unpack('<iii', doc.records[rec_idx]['payload'][:12])
        doc.records[rec_idx]['payload'] = bytearray(struct.pack('<iii', new_x, orig[1], orig[2]))
        doc.records[rec_idx]['size'] = 12

def update_t2150_w(doc, rec_idx, new_w):
    if rec_idx is not None:
        orig_flag = doc.records[rec_idx]['payload'][4:5]
        doc.records[rec_idx]['payload'] = bytearray(struct.pack('<i', new_w) + orig_flag)
        doc.records[rec_idx]['size'] = 5

def sync_and_save(doc, out_path, expected_recs=37198):
    for r in doc.records:
        r['size'] = len(r['payload'])
    assert len(doc.records) == expected_recs, f"Zero-shift violation! Expected {expected_recs}, got {len(doc.records)}"
    doc.save(out_path)
    print(f"   [SAVED] {os.path.basename(out_path)} ({len(doc.records):,} records) [PASS]")

def run_perfect_p1_p2():
    folder = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\ADHIKARYA PUTRA\New folder\JUN'
    base_xar = os.path.join(folder, '0.xar')
    excel_path = os.path.join(folder, 'Template_Pekerjaan_Xara_Jun.xlsx')

    print("=========================================================================")
    print("   EXECUTING PERFECT PIPELINE (PAGE 1 & 2): ADHIKARYA PUTRA JUN 2026")
    print("=========================================================================\n")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_mutasi = wb['Tabel_Mutasi']

    # Load 22 transactions for Page 1 & 2
    tx_list = []
    for r in range(6, 28): # Rows 1 to 22
        tx_list.append({
            'row_num': r - 5,
            'nominal': ws_mutasi.cell(r, 5).value,
            'balance': ws_mutasi.cell(r, 7).value
        })

    # Load Page 1 & Page 2 mapping
    with open('p1_p2_perfect_map.json', 'r', encoding='utf-8') as f:
        all_mapped_rows = json.load(f)

    # Filter out dummy rows (rows where saldo_rec is None)
    p1_p2_rows = [r for r in all_mapped_rows if r.get('saldo_rec') is not None]
    assert len(p1_p2_rows) == 22, f"Expected 22 valid rows, got {len(p1_p2_rows)}"
    print(f"[*] Loaded {len(p1_p2_rows)} mapped rows for Page 1 and Page 2.")

    # Base Document
    doc = XarDocument(base_xar)
    TOTAL_RECS = len(doc.records)

    # =========================================================================
    # [1] FIX HEADER NAMA & CABANG (PAGE 1 & 2)
    # =========================================================================
    print("\n[*] Applying Header Nama & Cabang fix (SOP Tahap 1: '\\r\\n' line-break)...")
    # Page 1: Rec 1010
    update_text(doc, 1010, "Adhikarya Putra \r\n")
    update_t2206_full(doc, 1009, 89858, 5761, 0)
    update_t2206_full(doc, 1014, 124101, 5761, 0)

    # Page 2: Rec 3645
    update_text(doc, 3645, "Adhikarya Putra \r\n")
    update_t2206_full(doc, 3644, 89858, 5761, 0)
    update_t2206_full(doc, 3649, 124101, 5761, 0)

    # =========================================================================
    # [2] FIX HEADER PERIODE (PAGE 1 & 2 CONTAINER EXPANSION)
    # =========================================================================
    print("[*] Applying Header Periode container expansion (Tag 2150 W=180.000 mp)...")
    # Page 1: Tag 2150 at Rec 1021 / 1039, Text at 1055, 1061
    update_t2150_w(doc, 1021, 180000)
    update_text(doc, 1055, "Jun 2026 - ")
    update_text(doc, 1061, "0 Jun 2026")

    # Page 2: Tag 2150 at Rec 3656 / 3674, Text at 3690, 3696
    update_t2150_w(doc, 3656, 180000)
    update_text(doc, 3690, "Jun 2026 - ")
    update_text(doc, 3696, "0 Jun 2026")

    # =========================================================================
    # [3] FIX HEADER DICETAK PADA (PAGE 1 & 2)
    # =========================================================================
    print("[*] Applying Header Dicetak Pada (10 Sep 2026)...")
    # Page 1: Rec 1073='1', 1077='0', 1085='Sep 2026'
    update_text(doc, 1073, "1")
    update_text(doc, 1077, "0")
    update_text(doc, 1085, "Sep 2026")

    # Page 2: Rec 3708='1', 3712='0', 3720='Sep 2026'
    update_text(doc, 3708, "1")
    update_text(doc, 3712, "0")
    update_text(doc, 3720, "Sep 2026")

    # =========================================================================
    # [4] FIX NOMOR REKENING (PAGE 1)
    # =========================================================================
    print("[*] Applying Header Nomor Rekening (1630016148929 )...")
    update_text(doc, 1108, "1630016148929 ")

    # =========================================================================
    # [5] FIX SUMMARY HEADER (CONTAINER EXPANSION & ALL 4 LINES)
    # =========================================================================
    print("[*] Applying Summary Header fix (Tag 2150 W=120.000 mp + 4 lines)...")
    update_t2150_w(doc, 1184, 120000) # Enlarge container so - 29.693.990,00 fits without wrapping!

    # Saldo Awal: Rec 1209
    str_sawal = "52.488,81 "
    update_text(doc, 1209, str_sawal)
    update_color(doc, 1205, COLOR_GRAY)
    update_t2206(doc, 1197, calc_text_width(str_sawal))

    # Dana Masuk: Rec 1219 ('+ 30.212.000,00'), Rec 1224 (blank)
    str_dmasuk = "+ 30.212.000,00"
    update_text(doc, 1219, str_dmasuk)
    blank_node(doc, 1224)
    update_color(doc, 1215, COLOR_GREEN)
    update_t2206(doc, 1213, calc_text_width(str_dmasuk))

    # Dana Keluar: Rec 1238 ('- 29.693.990,00 ')
    str_dkeluar = "- 29.693.990,00 "
    update_text(doc, 1238, str_dkeluar)
    update_color(doc, 1231, COLOR_BLACK)
    update_t2206(doc, 1229, calc_text_width(str_dkeluar))

    # Saldo Akhir: Rec 1252 ('570.498,81')
    str_sakhir = "570.498,81"
    update_text(doc, 1252, str_sakhir)
    update_color(doc, 1245, COLOR_BLUE)
    update_t2206(doc, 1242, calc_text_width(str_sakhir))

    # =========================================================================
    # [6] TAHAP 6: DATES & TIMES (PAGE 1 & 2)
    # =========================================================================
    print("[*] Applying Tahap 6 Dates & Times on Page 1 & 2...")
    for r_idx, m in enumerate(p1_p2_rows):
        # Update dates to Jun 2026
        for d_rec, d_txt in m['date_recs']:
            new_d = d_txt.replace('Apr 2026', 'Jun 2026').replace('Apr 202', 'Jun 202')
            update_text(doc, d_rec, new_d)

    # Save Tahap 6 checkpoint
    out_t6 = os.path.join(folder, '0_tahap6.xar')
    sync_and_save(doc, out_t6, TOTAL_RECS)

    # =========================================================================
    # [7] TAHAP 7: TABEL MUTASI (PAGE 1 & 2: ROWS 1 TO 22)
    # =========================================================================
    print("\n[*] Applying Tahap 7 Table Mutations on Page 1 & 2 (Rows 1 to 22)...")
    doc7 = XarDocument(out_t6)

    for r_idx, m in enumerate(p1_p2_rows):
        tx = tx_list[r_idx]
        row_num = r_idx + 1

        # 1. Update Saldo Text (Rec `saldo_rec`)
        s_rec = m['saldo_rec']
        s_splits = m.get('saldo_splits', [])
        bal_val = tx['balance']
        str_saldo = fmt_idr(bal_val)
        
        update_text(doc7, s_rec, str_saldo)
        for sp in s_splits:
            blank_node(doc7, sp)

        # 2. Update Nominal Text & Rata Kanan (Rec `nom_rec`, T2100 `nom_t2100`, T2206 `nom_t2206`)
        n_rec = m['nom_rec']
        n_splits = m.get('nom_splits', [])
        n_t2100 = m.get('nom_t2100')
        n_t2206 = m.get('nom_t2206')
        n_t150 = m.get('nom_t150')

        nom_val = tx['nominal']
        if nom_val > 0:
            str_nominal = f"+{fmt_idr(nom_val)}"
            nom_color = COLOR_GREEN
        else:
            str_nominal = f"-{fmt_idr(abs(nom_val))}"
            nom_color = COLOR_BLACK

        update_text(doc7, n_rec, str_nominal)
        for sp in n_splits:
            blank_node(doc7, sp)

        if n_t150:
            update_color(doc7, n_t150, nom_color)

        w_nominal = calc_text_width(str_nominal)
        new_x_left = TARGET_XR_NOMINAL - w_nominal
        
        if n_t2100:
            update_t2100_x(doc7, n_t2100, new_x_left)
        if n_t2206:
            update_t2206(doc7, n_t2206, w_nominal)

    out_t7 = os.path.join(folder, '0_tahap7.xar')
    sync_and_save(doc7, out_t7, TOTAL_RECS)

    print("\n=========================================================================")
    print("   [SUCCESS] FLALWESS PAGE 1 & PAGE 2 TAHAP 7 GENERATED!")
    print(f"   Saved Output: {out_t7}")
    print("=========================================================================\n")

if __name__ == '__main__':
    run_perfect_p1_p2()
