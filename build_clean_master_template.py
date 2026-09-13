import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_pristine_template(output_path, fill_july_data=True):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    dark_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    section_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    light_blue_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    green_soft = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    font_title = Font(name="Segoe UI", size=12, bold=True, color="1F4E79")
    font_sub = Font(name="Segoe UI", size=9, italic=True, color="595959")
    font_section = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_status_wajib = Font(name="Segoe UI", size=9, bold=True, color="C00000")
    font_status_ops = Font(name="Segoe UI", size=9, italic=True, color="7F7F7F")
    font_status_auto = Font(name="Segoe UI", size=9, bold=True, color="0070C0")

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color="1B365D"))

    # =========================================================================
    # SHEET 1: Header & Ringkasan (ZERO MERGED CELLS)
    # =========================================================================
    ws1 = wb.create_sheet(title="Header & Ringkasan")
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "KONFIGURASI UTAMA PEKERJAAN (TAHAP 1 - 5 & RINGKASAN TAHAP 7)"
    ws1["A1"].font = font_title
    ws1["A2"] = "Template resmi Xara Project V2 (Bebas Merged Cells - Aman untuk seleksi, pengosongan sel, & copy-paste)"
    ws1["A2"].font = font_sub

    headers1 = ["Kategori / Parameter", "Nilai Data (Input Anda)", "Status", "Keterangan & Aturan Kerja (SOP Project V2)"]
    for c_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = dark_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_header
    ws1.row_dimensions[4].height = 26

    # Sections definition
    sections = [
        ("1. INFORMASI FILE & PATH PROYEK", [
            ("File Sumber (.xar)", r"C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Firmansyah\New folder\Jul\0.xar", "Wajib", "Lokasi file .xar template asli yang akan diedit"),
            ("File Output (.xar)", r"C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Firmansyah\New folder\Jul\0_output.xar", "Wajib", "Lokasi file .xar baru hasil editan (file asli tetap aman)"),
            ("Nama Dokumen / Catatan", "Rekening Koran JUL 2026 - FIRMANSYAH", "Opsional", "Keterangan label proyek untuk arsip riwayat training")
        ]),
        ("2. DATA HEADER NASABAH (TAHAP 1 s.d. 4)", [
            ("Nama Nasabah (Tahap 1)", "FIRMANSYAH", "Wajib", "Mengubah Nama/Name di semua lembar header secara otomatis"),
            ("Periode Laporan (Tahap 2)", "01 Jul 2026 - 31 Jul 2026", "Wajib", "Mengubah rentang periode di seluruh halaman header"),
            ("Dicetak Pada (Tahap 3)", "10 Sep 2026", "Wajib", "Mengubah tanggal cetak / Issued on di seluruh lembar halaman"),
            ("Nomor Rekening (Tahap 4)", "1630010942426", "Wajib", "Nomor rekening nasabah (format teks agar tidak menjadi eksponen/E+)")
        ]),
        ("3. PENOMORAN HALAMAN (TAHAP 5)", [
            ("Mode Total Halaman", "AUTO", "Direkomendasikan", "Ketik 'AUTO' agar sistem mendeteksi total lembar fisik secara mandiri"),
            ("Total Halaman Manual", "3", "Opsional", "Hanya digunakan jika ingin memaksa angka total halaman tertentu")
        ]),
        ("4. BULAN & TAHUN TRANSAKSI (TAHAP 6)", [
            ("Target Bulan & Tahun", "Jul 2026", "Wajib", "Bulan dan tahun transaksi berjalan"),
            ("Digit Penutup Tahun", "6", "Wajib", "Mengganti node digit penutup tahun (misal tahun 2026 -> digit 6)")
        ]),
        ("5. RINGKASAN KEUANGAN HEADER (TAHAP 7)", [
            ("Saldo Awal", 1929504, "Wajib", "Saldo awal rekening awal periode (menjadi saldo pembuka di Tabel Mutasi)"),
            ("Dana Masuk (Kredit)", 9197167, "Wajib", "Total mutasi uang masuk pada periode ini (Warna Hijau #ba030000)"),
            ("Dana Keluar (Debit)", 5234704, "Wajib", "Total mutasi uang keluar pada periode ini (Warna Hitam #87010000)"),
            ("Saldo Akhir", 5891967, "Wajib", "Saldo penutupan akhir periode (Warna Biru #0d050000)"),
            ("Keseimbangan Neraca (Audit)", '=IF(ROUND(B21+B22-B23-B24,2)=0,"BALANCE (MATCH)","SELISIH: " & TEXT(B21+B22-B23-B24,"#,##0.00"))', "Auto-Check", "Audit Live Formula: Saldo Awal (B21) + Masuk (B22) - Keluar (B23) = Saldo Akhir (B24)")
        ])
    ]

    r_idx = 5
    for sec_title, items in sections:
        # Section banner (styled individual cells, NO MERGE!)
        for c in range(1, 5):
            cell = ws1.cell(row=r_idx, column=c)
            cell.fill = section_fill
            cell.border = border_data
            if c == 1:
                cell.value = sec_title
                cell.font = font_section
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws1.row_dimensions[r_idx].height = 22
        r_idx += 1

        for p_name, p_val, status, ket in items:
            c1 = ws1.cell(row=r_idx, column=1, value=p_name)
            c2 = ws1.cell(row=r_idx, column=2, value=p_val)
            c3 = ws1.cell(row=r_idx, column=3, value=status)
            c4 = ws1.cell(row=r_idx, column=4, value=ket)

            c1.font = font_body
            c2.font = font_bold if "Ringkasan" in sec_title or status == "Wajib" else font_body
            c3.font = font_status_wajib if status == "Wajib" else (font_status_auto if "Auto" in status or "Rekomendasi" in status else font_status_ops)
            c4.font = font_sub

            c1.border = border_data
            c2.border = border_data
            c3.border = border_data
            c4.border = border_data

            c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            c3.alignment = Alignment(horizontal="center", vertical="center")
            c4.alignment = Alignment(horizontal="left", vertical="center")

            # Format formatting for column B
            if isinstance(p_val, (int, float)):
                c2.number_format = "#,##0.00"
                c2.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c2.number_format = "@"
                c2.alignment = Alignment(horizontal="left", vertical="center")

            if p_name == "Keseimbangan Neraca (Audit)":
                c2.fill = green_soft
                c2.alignment = Alignment(horizontal="center", vertical="center")

            ws1.row_dimensions[r_idx].height = 20
            r_idx += 1

    ws1.column_dimensions['A'].width = 32
    ws1.column_dimensions['B'].width = 35
    ws1.column_dimensions['C'].width = 18
    ws1.column_dimensions['D'].width = 65


    # =========================================================================
    # SHEET 2: Tabel_Mutasi (ZERO MERGED CELLS & CLEAN)
    # =========================================================================
    ws2 = wb.create_sheet(title="Tabel_Mutasi")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "TABEL MUTASI TRANSAKSI UTAMA"
    ws2["A1"].font = font_title
    ws2["A2"] = "Kolom Tanggal dapat dikosongkan pada baris transaksi yang berada di hari yang sama. Kolom No wajib berupa angka urut."
    ws2["A2"].font = font_sub

    headers2 = ["No", "Tanggal", "Jam", "Uraian Transaksi", "Nominal (+/-)", "Tipe (CR/DB)", "Saldo Berjalan", "Kustomisasi Warna", "Kustomisasi Alignment", "Catatan Baris"]
    for c_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = dark_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_header
    ws2.row_dimensions[4].height = 26

    # Row 5: Saldo Awal (ZERO MERGE)
    r5_vals = ["[AWAL]", "01/07", "-", "SALDO AWAL PEMBUKAAN PERIODE (INITIAL BALANCE)", "-", "AWAL", "='Header & Ringkasan'!B21", "AUTO (Tag 150)", "AUTO (20.049 cm)", "Saldo awal pembukaan rekening periode laporan"]
    for c_idx, val in enumerate(r5_vals, 1):
        cell = ws2.cell(row=5, column=c_idx, value=val)
        cell.font = font_bold
        cell.fill = light_blue_fill
        cell.border = border_data
        if c_idx in (1, 2, 3, 6):
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.number_format = "@"
        elif c_idx == 7:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0.00"
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.number_format = "@"
    ws2.row_dimensions[5].height = 20

    # 22 Active Transactions for July (Excluding zero rows 11 and 12)
    # Date formatting follows user's natural format: populated on date change, blank on repeat!
    tx_raw = [
        (1,  "01/07", "04:00:00 WIB", "-", -397500,  "DB", 1532004),
        (2,  "",      "04:12:15 WIB", "-", -300000,  "DB", 1232004),
        (3,  "",      "08:30:00 WIB", "-", -42000,   "DB", 1190004),
        (4,  "02/07", "09:15:22 WIB", "-", -500000,  "DB", 690004),
        (5,  "",      "10:05:10 WIB", "-", 15200,    "CR", 705204),
        (6,  "",      "11:20:45 WIB", "-", -8000,    "DB", 697204),
        (7,  "",      "12:00:00 WIB", "-", -18000,   "DB", 679204),
        (8,  "",      "13:45:12 WIB", "-", 500000,   "CR", 1179204),
        (9,  "05/07", "14:10:00 WIB", "-", -500704,  "DB", 678500),
        (10, "06/07", "15:25:30 WIB", "-", 26967,    "CR", 705467),
        (11, "08/07", "09:30:00 WIB", "-", 500000,   "CR", 1205467),
        (12, "09/07", "10:00:00 WIB", "-", -350500,  "DB", 854967),
        (13, "10/07", "11:15:20 WIB", "-", -130000,  "DB", 724967),
        (14, "",      "12:35:40 WIB", "-", -45000,   "DB", 679967),
        (15, "",      "13:10:00 WIB", "-", 500000,   "CR", 1179967),
        (16, "",      "04:00:00 WIB", "-", 7605000,  "CR", 8784967),
        (17, "",      "15:05:00 WIB", "-", -450000,  "DB", 8334967),
        (18, "14/07", "09:00:00 WIB", "-", -1500000, "DB", 6834967),
        (19, "15/07", "10:45:10 WIB", "-", -500000,  "DB", 6334967),
        (20, "",      "11:12:00 WIB", "-", 50000,    "CR", 6384967),
        (21, "16/07", "12:00:15 WIB", "-", -488000,  "DB", 5896967),
        (22, "31/07", "23:59:00 WIB", "-", -5000,    "DB", 5891967)
    ]

    r_curr = 6
    for no, tgl, jam, uraian, nom, tipe, saldo in tx_raw:
        row_vals = [
            no, tgl, jam, uraian, nom, tipe, saldo,
            "AUTO (Tag 150)", "AUTO (15.214 cm)", "OK Verified"
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws2.cell(row=r_curr, column=c_idx, value=val)
            cell.font = font_body
            cell.border = border_data
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "0"
            elif c_idx in (2, 3, 4, 6, 8, 9, 10):
                cell.alignment = Alignment(horizontal="center" if c_idx in (2, 3, 6) else "left", vertical="center")
                cell.number_format = "@"
            elif c_idx in (5, 7):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
                if c_idx == 5:
                    if isinstance(val, (int, float)) and val > 0:
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="008000")
        ws2.row_dimensions[r_curr].height = 19
        r_curr += 1

    # Optional blank rows for user expansion (Clean rows 28 s.d. 35)
    for blank_r in range(r_curr, r_curr + 8):
        for c_idx in range(1, 11):
            cell = ws2.cell(row=blank_r, column=c_idx)
            cell.border = border_data
            cell.font = font_body
            if c_idx == 1:
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in (5, 7):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in (8, 9, 10):
                cell.value = "AUTO (Tag 150)" if c_idx == 8 else ("AUTO (15.214 cm)" if c_idx == 9 else "OK Verified")
                cell.number_format = "@"
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.number_format = "@"
                cell.alignment = Alignment(horizontal="center" if c_idx in (2, 3, 6) else "left", vertical="center")
        ws2.row_dimensions[blank_r].height = 19

    ws2.column_dimensions['A'].width = 8
    ws2.column_dimensions['B'].width = 14
    ws2.column_dimensions['C'].width = 16
    ws2.column_dimensions['D'].width = 32
    ws2.column_dimensions['E'].width = 18
    ws2.column_dimensions['F'].width = 14
    ws2.column_dimensions['G'].width = 18
    ws2.column_dimensions['H'].width = 18
    ws2.column_dimensions['I'].width = 20
    ws2.column_dimensions['J'].width = 18


    # =========================================================================
    # SHEET 3: Penyesuaian_Tambahan (ZERO MERGED CELLS & CLEAN)
    # =========================================================================
    ws3 = wb.create_sheet(title="Penyesuaian_Tambahan")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = "LEMBAR KHUSUS PENYESUAIAN TAMBAHAN (ADVANCED CUSTOMIZATION)"
    ws3["A1"].font = font_title
    ws3["A2"] = "Konfigurasi tingkat lanjut untuk ruler koordinat, palet warna biner, dan proteksi anti-corruption."
    ws3["A2"].font = font_sub

    headers3 = ["Kategori Penyesuaian", "Nama Parameter / Teks Asal", "Nilai Kustom / Teks Baru", "Target Halaman / Elemen", "Keterangan & Panduan"]
    for c_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = dark_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_header
    ws3.row_dimensions[4].height = 26

    sections3 = [
        ("A. KUSTOMISASI RULER & BATAS RATA KANAN (RIGHT-ALIGNMENT GRID)", [
            ("Ruler Grid", "Target X-Right Kolom Nominal", "431250 mp (15.214 cm)", "Semua Baris Nominal", "Titik koordinat batas sisi kanan kolom Nominal pada Xara"),
            ("Ruler Grid", "Target X-Right Kolom Saldo", "570250 mp (20.049 cm)", "Semua Baris Saldo", "Titik koordinat batas sisi kanan kolom Saldo pada Xara"),
            ("Ruler Grid", "Mode Kalkulasi Otomatis Width", "AKTIF", "Tag 2206 & 2100", "Menghitung mundur X_left = X_target - W(teks) sesuai glyph metrics font")
        ]),
        ("B. KUSTOMISASI PALET WARNA BINER (TAG 150)", [
            ("Warna Biner", "Nominal Kredit (CR)", "b'\\xba\\x03\\x00\\x00' (#00A651)", "Tabel Mutasi", "Kunci warna Hijau resmi bank Mandiri"),
            ("Warna Biner", "Nominal Debit (DB)", "b'\\x87\\x01\\x00\\x00' (#000000)", "Tabel Mutasi", "Kunci warna Hitam resmi bank Mandiri"),
            ("Warna Biner", "Saldo Berjalan", "b'\\x0d\\x05\\x00\\x00' (#005B9C)", "Tabel Mutasi", "Kunci warna Biru resmi bank Mandiri"),
            ("Warna Biner", "Header Dana Masuk", "b'\\xba\\x03\\x00\\x00' (#00A651)", "Header Ringkasan", "Kunci warna Hijau ringkasan"),
            ("Warna Biner", "Header Dana Keluar", "b'\\x87\\x01\\x00\\x00' (#000000)", "Header Ringkasan", "Kunci warna Hitam ringkasan"),
            ("Warna Biner", "Header Saldo Akhir", "b'\\x0d\\x05\\x00\\x00' (#005B9C)", "Header Ringkasan", "Kunci warna Biru ringkasan")
        ]),
        ("C. OVERRIDE TEKS KHUSUS / ALAMAT / CABANG TAMBAHAN", [
            ("Override Teks", "KCP JAKARTA SUDIRMAN", "KCP JAKARTA SUDIRMAN", "Header Lembar 1", "Dapat diubah jika ingin mengganti nama kantor cabang"),
            ("Override Teks", "JL. JEND. SUDIRMAN KAV. 52-53", "JL. JEND. SUDIRMAN KAV. 52-53", "Header Lembar 1", "Dapat diubah jika ingin mengganti alamat nasabah"),
            ("Override Teks", "JAKARTA SELATAN 12190", "JAKARTA SELATAN 12190", "Header Lembar 1", "Dapat diubah jika ingin mengganti kota dan kode pos"),
            ("Override Teks", "IDR", "IDR", "Semua Halaman", "Mata uang laporan rekening"),
            ("Override Teks", "[Teks Kustom Lain]", "[Pengganti Kustom]", "[Halaman Target]", "Slot bebas untuk teks tambahan lain yang ingin diganti")
        ]),
        ("D. FITUR KEAMANAN BINER & ANTI-CORRUPTION (STABILITY SWITCHES)", [
            ("Keamanan Biner", "Zero-Shift Record Invariant", "TERKUNCI (AKTIF)", "Seluruh Stream", "Mencegah penambahan/pengurangan record agar ID warna & font tidak rusak"),
            ("Keamanan Biner", "Tag 2202 Null Character Safety", "b'\\x00\\x00' (2 Bytes)", "Pecahan Node", "Menjamin node sekunder tidak pernah 0 byte agar Xara tidak crash"),
            ("Keamanan Biner", "Auto-Sync Payload Size", "AKTIF", "Header Tiap Record", "rec['size'] = len(rec['payload']) selalu sinkron sebelum file disimpan"),
            ("Keamanan Biner", "Lock Embedded Font (Tag 2907)", "TERKUNCI (AKTIF)", "Font Definitions", "Mencegah font fallback ter-reset menjadi default PDF")
        ]),
        ("E. LOG CATATAN OPERATOR / INSTRUKSI KHUSUS PEKERJAAN", [
            ("Catatan Operator", "Catatan Klien / Khusus", "Tidak ada instruksi khusus", "Audit Trail", "Catatan bebas untuk operator sebelum dokumen diekspor")
        ])
    ]

    r3_idx = 5
    for sec_title, items in sections3:
        for c in range(1, 6):
            cell = ws3.cell(row=r3_idx, column=c)
            cell.fill = section_fill
            cell.border = border_data
            if c == 1:
                cell.value = sec_title
                cell.font = font_section
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws3.row_dimensions[r3_idx].height = 22
        r3_idx += 1

        for kat, param, val, target, ket in items:
            c1 = ws3.cell(row=r3_idx, column=1, value=kat)
            c2 = ws3.cell(row=r3_idx, column=2, value=param)
            c3 = ws3.cell(row=r3_idx, column=3, value=val)
            c4 = ws3.cell(row=r3_idx, column=4, value=target)
            c5 = ws3.cell(row=r3_idx, column=5, value=ket)

            c1.font = font_body
            c2.font = font_body
            c3.font = font_bold
            c4.font = font_body
            c5.font = font_sub

            for cell in (c1, c2, c3, c4, c5):
                cell.border = border_data
                cell.number_format = "@"

            c1.alignment = Alignment(horizontal="left", vertical="center")
            c2.alignment = Alignment(horizontal="left", vertical="center")
            c3.alignment = Alignment(horizontal="center" if "AKTIF" in str(val) or "TERKUNCI" in str(val) else "left", vertical="center")
            c4.alignment = Alignment(horizontal="left", vertical="center")
            c5.alignment = Alignment(horizontal="left", vertical="center")

            ws3.row_dimensions[r3_idx].height = 20
            r3_idx += 1

    ws3.column_dimensions['A'].width = 22
    ws3.column_dimensions['B'].width = 32
    ws3.column_dimensions['C'].width = 32
    ws3.column_dimensions['D'].width = 25
    ws3.column_dimensions['E'].width = 65

    wb.save(output_path)
    print(f"[OK] Pristine Bug-Free Template saved to: {output_path}")

if __name__ == '__main__':
    # 1. Update Master in xara_copilot
    p_master = r'C:\Users\Lenovo\xara_copilot\Template_Pekerjaan_Xara.xlsx'
    build_pristine_template(p_master)

    # 2. Update July working template
    p_jul = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Firmansyah\New folder\Jul\Template_Pekerjaan_Xara_JUL.xlsx'
    p_jul_clean = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Firmansyah\New folder\Jul\Template_Pekerjaan_Xara_JUL_Clean.xlsx'
    try:
        build_pristine_template(p_jul)
    except PermissionError:
        print(f"[NOTE] {p_jul} sedang dibuka di Microsoft Excel. Menyimpan ke: {p_jul_clean}")
        build_pristine_template(p_jul_clean)

