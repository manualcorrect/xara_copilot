import os
import openpyxl
import json
import random
from datetime import datetime, time
from xar_dom_engine import XarDocument

# 1. Base files
folder = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\aug'
excel_path = os.path.join(folder, 'Template_Pekerjaan_Xara_Aug.xlsx')
xar_path = os.path.join(folder, '0.xar')

doc = XarDocument(xar_path)
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws_mutasi = wb['Tabel_Mutasi']

# 2. Extract transaction amounts and user dates/times from Excel
tx_data = []
for r in range(6, 116):
    tx_idx = r - 5 # 1 to 110
    no_val = ws_mutasi.cell(r, 1).value
    tgl_val = ws_mutasi.cell(r, 2).value
    jam_val = ws_mutasi.cell(r, 3).value
    uraian_val = ws_mutasi.cell(r, 4).value
    nom_val = ws_mutasi.cell(r, 5).value
    tipe_val = ws_mutasi.cell(r, 6).value
    saldo_val = ws_mutasi.cell(r, 7).value
    
    nom_float = float(nom_val) if nom_val is not None else 0.0
    saldo_float = float(saldo_val) if saldo_val is not None else 0.0
    
    # User explicit date/time
    explicit_date = None
    explicit_time = None
    
    # Check if tgl_val is datetime or valid date
    if isinstance(tgl_val, datetime):
        explicit_date = f"{tgl_val.day:02d} Aug 2026"
    elif isinstance(tgl_val, str) and '/' in tgl_val and not any(w in tgl_val for w in ['Users', 'BENDI', 'AUG', 'Rekening']):
        parts = tgl_val.strip().split('/')
        try:
            day = int(parts[0])
            explicit_date = f"{day:02d} Aug 2026"
        except:
            pass
            
    # Check if jam_val is time or valid time string
    if isinstance(jam_val, time):
        explicit_time = f"{jam_val.strftime('%H:%M:%S')} WIB"
    elif isinstance(jam_val, str) and (':' in jam_val or 'WIB' in jam_val) and not any(w in jam_val for w in ['Users', 'BENDI', 'AUG', 'Rekening']):
        s = jam_val.strip()
        if 'WIB' in s:
            explicit_time = s
        else:
            explicit_time = f"{s} WIB"
            
    tx_data.append({
        'tx_idx': tx_idx,
        'excel_row': r,
        'explicit_date': explicit_date,
        'explicit_time': explicit_time,
        'nom_float': nom_float,
        'saldo_float': saldo_float,
        'uraian': uraian_val
    })

print(f"Loaded {len(tx_data)} rows from Excel")
for t in tx_data:
    if t['explicit_date'] or t['explicit_time']:
        print(f"  Tx {t['tx_idx']:3d} (Excel row {t['excel_row']:3d}): Date={t['explicit_date']}, Time={t['explicit_time']}, Nom={t['nom_float']}, Saldo={t['saldo_float']}")
