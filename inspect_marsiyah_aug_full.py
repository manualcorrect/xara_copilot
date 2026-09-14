import openpyxl
from xar_dom_engine import XarDocument
import struct

folder = r'C:\Users\Lenovo\Downloads\rekening\PT BENDI NASHA NIAGA INDUSTRI\Marsiyah\New folder\aug'
excel_path = f"{folder}\\Template_Pekerjaan_Xara_Aug.xlsx"
xar_path = f"{folder}\\0.xar"

print("=========================================================================")
print("   INSPECTING MARSIYAH AGUSTUS 2026 WORKSPACE")
print("=========================================================================\n")

# 1. Parse Excel
wb = openpyxl.load_workbook(excel_path, data_only=True)
print(f"[*] Sheet names: {wb.sheetnames}")

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, min(20, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        if any(v is not None for v in row_vals):
            print(f"  Row {r}: {row_vals}")

# 2. Parse 0.xar
doc = XarDocument(xar_path)
print(f"\n[*] 0.xar Total Records: {len(doc.records):,}")

# Find pages in 0.xar
page_tags = [i for i, r in enumerate(doc.records) if r['tag'] == 4351]
print(f"[*] Found {len(page_tags)} page markers (Tag 4351): {page_tags}")

# Find all occurrences of Page numbers (e.g. '1 dari ...' or '1 of ...')
for i, r in enumerate(doc.records):
    if r['tag'] in (2201, 2202):
        txt = r['payload'].decode('utf-16le', errors='ignore')
        if 'dari' in txt or ' of ' in txt:
            print(f"  Page Num Record [{i}] Tag {r['tag']}: {txt!r}")

# Check Font Definitions and Colors in 0.xar
print("\n=== FONT DEFINITIONS IN 0.xar ===")
for i, r in enumerate(doc.records[:500]):
    if r['tag'] in (2900, 2905, 2906, 2907, 2908, 2909, 2910):
        print(f"  Rec [{i}] Tag {r['tag']} (sz {r['size']}): {r['payload'].hex()}")

print("\n=== NAME & CABANG IN 0.xar ===")
for i, r in enumerate(doc.records[:2000]):
    if r['tag'] in (2201, 2202):
        txt = r['payload'].decode('utf-16le', errors='ignore')
        if any(w in txt for w in ['ANWAR', 'DINI', 'MASRIYAH', 'MOHAMMAD', 'ROY', 'KCP', 'Jakarta', 'Aries', 'Nama', 'Cabang']):
            print(f"  Rec [{i}] Tag {r['tag']}: {txt!r}")
